#!/usr/bin/env python3
"""Create a group-meeting deck section for Visium HD Exp1 training results."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path("/Users/haoranwu/Desktop/Yan_Lab_Research/Medical-SAM3")
SOURCE_DECK = Path("/Users/haoranwu/Desktop/Yan_Lab_Research/Presentation/Apr30.pptx")
OUT_DIR = ROOT / "inference" / "visium_hd_exp1"
FINAL = OUT_DIR / "Medical_SAM3_Training_Ablation_GroupMeeting_20260507.pptx"
FIG_DIR = ROOT / "results" / "visium_hd_exp1" / "figures" / "group_meeting"
APR30_RENDER_DIR = ROOT / "output" / "group_meeting_ppt" / "render_check_20260507"


NAVY = RGBColor(18, 31, 49)
TEAL = RGBColor(24, 121, 113)
BLUE = RGBColor(37, 99, 235)
AMBER = RGBColor(180, 83, 9)
GREEN = RGBColor(22, 101, 52)
RED = RGBColor(185, 28, 28)
GRAY = RGBColor(100, 116, 139)
LIGHT = RGBColor(248, 250, 252)
LINE = RGBColor(203, 213, 225)
WHITE = RGBColor(255, 255, 255)


def delete_slides_from(prs: Presentation, start_index_zero_based: int) -> None:
    slide_ids = prs.slides._sldIdLst  # noqa: SLF001 - python-pptx has no public delete API.
    while len(prs.slides) > start_index_zero_based:
        slide_ids.remove(slide_ids[start_index_zero_based])


def add_apr30_context_slides(prs: Presentation) -> None:
    """Use rendered images for the old context slides to avoid fragile PPTX copying."""
    for i in range(1, 8):
        slide = blank(prs)
        image_path = APR30_RENDER_DIR / f"apr30_slide-{i:02d}.png"
        add_picture(slide, image_path, Inches(0), Inches(0), w=Inches(13.333), h=Inches(7.5))


def add_future_directions_overview(prs: Presentation) -> None:
    """Recreate the one prior slide the user wants to keep as the bridge."""
    slide = blank(prs)
    # Top banner matching the user's prior Future Directions slide.
    banner = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.05))
    banner.fill.solid()
    banner.fill.fore_color.rgb = RGBColor(35, 59, 103)
    banner.line.fill.background()
    add_text(slide, Inches(0.42), Inches(0.20), Inches(7.7), Inches(0.36), "Future Directions: Overview", 25, WHITE, True)
    add_text(slide, Inches(0.42), Inches(0.72), Inches(5.0), Inches(0.26), "Five Ways to Intervene in the SAM3 Pipeline", 13, RGBColor(226, 232, 240), True)
    add_text(slide, Inches(3.92), Inches(1.18), Inches(5.6), Inches(0.28), "SAM3 Architecture - Where Each Future Direction Intervenes", 14, NAVY, True, PP_ALIGN.CENTER)

    # Main SAM3 path.
    y = Inches(3.45)
    add_label(slide, Inches(0.45), y, Inches(2.05), Inches(0.86), "H&E Image\n(tile 1024px)", GRAY, RGBColor(229, 231, 235), 12)
    add_label(slide, Inches(3.0), y, Inches(2.4), Inches(0.86), "Image\nEncoder (ViT)", RGBColor(30, 64, 120), RGBColor(219, 234, 254), 13)
    add_label(slide, Inches(6.05), y, Inches(2.3), Inches(0.86), "Prompt\nEncoder", RGBColor(30, 64, 120), RGBColor(219, 234, 254), 13)
    add_label(slide, Inches(9.05), y, Inches(2.25), Inches(0.86), "Mask\nDecoder", RGBColor(30, 64, 120), RGBColor(219, 234, 254), 13)
    add_label(slide, Inches(11.78), y, Inches(1.25), Inches(0.86), "Mask\nOutput", BLUE, RGBColor(219, 234, 254), 13)
    for x1, x2 in [(2.5, 3.0), (5.4, 6.05), (8.35, 9.05), (11.3, 11.78)]:
        line = slide.shapes.add_connector(1, Inches(x1), Inches(3.88), Inches(x2), Inches(3.88))
        line.line.color.rgb = GRAY
        line.line.width = Pt(2)

    # Intervention labels.
    add_text(slide, Inches(3.45), Inches(3.05), Inches(2.2), Inches(0.28), "B. GigaPath Backbone", 10, BLUE, True, PP_ALIGN.CENTER)
    add_text(slide, Inches(3.35), Inches(3.26), Inches(2.4), Inches(0.28), "A. Supervised Fine-tune", 10, AMBER, True, PP_ALIGN.CENTER)
    add_text(slide, Inches(11.85), Inches(4.28), Inches(1.25), Inches(0.38), "E. Replace with\nCellViT / HoVerNet", 9, RED, True, PP_ALIGN.CENTER)

    # Expression path and RAG branch.
    add_label(slide, Inches(1.0), Inches(2.05), Inches(2.25), Inches(0.68), "Gene Expression\n(18K genes)", GREEN, RGBColor(220, 252, 231), 11)
    add_label(slide, Inches(5.95), Inches(2.05), Inches(2.75), Inches(0.68), "Expression\nMLP Encoder", GREEN, RGBColor(220, 252, 231), 12)
    line = slide.shapes.add_connector(1, Inches(3.55), Inches(2.39), Inches(5.95), Inches(2.39))
    line.line.color.rgb = GREEN
    line.line.width = Pt(1.8)
    line = slide.shapes.add_connector(1, Inches(7.33), Inches(2.73), Inches(7.33), Inches(3.45))
    line.line.color.rgb = GREEN
    line.line.width = Pt(2)
    add_text(slide, Inches(6.78), Inches(2.76), Inches(1.4), Inches(0.42), "C. Expression\n-> Prompt", 10, GREEN, True, PP_ALIGN.CENTER)

    add_label(slide, Inches(7.35), Inches(1.48), Inches(3.55), Inches(0.68), "Reference DB\n(annotated patches)", RGBColor(126, 34, 206), RGBColor(243, 232, 255), 11)
    line = slide.shapes.add_connector(1, Inches(9.0), Inches(3.45), Inches(9.0), Inches(2.16))
    line.line.color.rgb = RGBColor(126, 34, 206)
    line.line.width = Pt(1.8)
    add_text(slide, Inches(10.95), Inches(1.42), Inches(2.1), Inches(0.45), "D. RAG: retrieve similar\nannotated patches", 10, RGBColor(126, 34, 206), True)

    # Existing prompt modalities.
    add_label(slide, Inches(5.48), Inches(4.78), Inches(1.15), Inches(0.55), "Points\n/ Boxes", GRAY, RGBColor(229, 231, 235), 9)
    add_label(slide, Inches(6.66), Inches(4.78), Inches(1.05), Inches(0.55), "Text\n(CLIP)", GRAY, RGBColor(229, 231, 235), 9)
    add_label(slide, Inches(7.74), Inches(4.78), Inches(1.25), Inches(0.55), "Exemplar\nPatches", GRAY, RGBColor(229, 231, 235), 9)
    for x in [6.05, 7.18, 8.35]:
        line = slide.shapes.add_connector(1, Inches(x), Inches(4.78), Inches(7.20), Inches(4.31))
        line.line.color.rgb = GRAY
        line.line.width = Pt(1.2)

    # Legend.
    rows = [
        ["A", "Supervised fine-tune image encoder + mask decoder"],
        ["B", "Replace image encoder with pathology foundation model"],
        ["C", "Expression MLP -> Prompt Encoder"],
        ["D", "RAG: retrieval exemplars -> Prompt Encoder"],
        ["E", "Replace whole SAM3 with pathology model"],
    ]
    add_table(slide, Inches(0.28), Inches(5.85), Inches(3.6), Inches(1.25), ["", "Direction"], rows, [Inches(0.35), Inches(3.25)], 7.1)



def blank(prs: Presentation):
    return prs.slides.add_slide(prs.slide_layouts[6])


def add_title(slide, title: str, subtitle: str | None = None) -> None:
    box = slide.shapes.add_textbox(Inches(0.55), Inches(0.28), Inches(12.25), Inches(0.55))
    p = box.text_frame.paragraphs[0]
    p.text = title
    p.font.name = "Aptos Display"
    p.font.size = Pt(26)
    p.font.bold = True
    p.font.color.rgb = NAVY
    if subtitle:
        sub = slide.shapes.add_textbox(Inches(0.58), Inches(0.82), Inches(11.8), Inches(0.34))
        q = sub.text_frame.paragraphs[0]
        q.text = subtitle
        q.font.name = "Aptos"
        q.font.size = Pt(10.5)
        q.font.color.rgb = GRAY
    slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(0.55),
        Inches(1.18),
        Inches(12.2),
        Inches(0.02),
    ).fill.solid()
    slide.shapes[-1].fill.fore_color.rgb = LINE
    slide.shapes[-1].line.fill.background()


def add_footer(slide, text: str = "Haoran Wu | Group meeting | May 2026") -> None:
    box = slide.shapes.add_textbox(Inches(0.55), Inches(7.08), Inches(12.2), Inches(0.22))
    p = box.text_frame.paragraphs[0]
    p.text = text
    p.font.name = "Aptos"
    p.font.size = Pt(8)
    p.font.color.rgb = RGBColor(148, 163, 184)


def add_label(slide, x, y, w, h, text, color=TEAL, fill=LIGHT, size=12, bold=True):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = color
    shape.line.width = Pt(1.2)
    tf = shape.text_frame
    tf.clear()
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.08)
    tf.margin_top = Inches(0.04)
    tf.margin_bottom = Inches(0.04)
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = PP_ALIGN.CENTER
    p.font.name = "Aptos"
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    return shape


def add_text(slide, x, y, w, h, text, size=12, color=NAVY, bold=False, align=None):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = "Aptos"
    p.font.size = Pt(size)
    p.font.color.rgb = color
    p.font.bold = bold
    if align is not None:
        p.alignment = align
    return box


def add_metric(slide, x, y, w, h, value, label, color):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    tf = shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = value
    p.alignment = PP_ALIGN.CENTER
    p.font.name = "Aptos Display"
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = WHITE
    q = tf.add_paragraph()
    q.text = label
    q.alignment = PP_ALIGN.CENTER
    q.font.name = "Aptos"
    q.font.size = Pt(8.5)
    q.font.color.rgb = WHITE


def add_table(slide, x, y, w, h, headers, rows, col_widths=None, font_size=8.8):
    table = slide.shapes.add_table(len(rows) + 1, len(headers), x, y, w, h).table
    if col_widths:
        for idx, width in enumerate(col_widths):
            table.columns[idx].width = width
    for c, header in enumerate(headers):
        cell = table.cell(0, c)
        cell.text = header
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
        for p in cell.text_frame.paragraphs:
            p.font.name = "Aptos"
            p.font.size = Pt(font_size)
            p.font.bold = True
            p.font.color.rgb = WHITE
    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            cell.text = str(val)
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if r % 2 else RGBColor(241, 245, 249)
            for p in cell.text_frame.paragraphs:
                p.font.name = "Aptos"
                p.font.size = Pt(font_size)
                p.font.color.rgb = NAVY
    return table


def add_picture(slide, path: Path, x, y, w=None, h=None):
    if not path.exists():
        add_text(slide, x, y, w or Inches(5), h or Inches(0.5), f"Missing figure:\n{path}", size=9, color=RED)
        return None
    return slide.shapes.add_picture(str(path), x, y, width=w, height=h)


def slide_training_overview(prs):
    slide = blank(prs)
    add_title(slide, "What I trained after the SAM3 baseline", "Goal: learn a shared H&E-expression representation before using expression to guide prompting/retrieval")
    add_footer(slide)
    add_text(slide, Inches(0.72), Inches(1.48), Inches(5.15), Inches(0.38), "Problem after Apr30", 15, NAVY, True)
    add_text(
        slide,
        Inches(0.72),
        Inches(1.92),
        Inches(5.1),
        Inches(1.25),
        "Oracle SAM3 prompts recover semantic tissue regions only weakly. Instead of asking SAM3 to solve the whole task directly, I trained a cross-modal model that connects local H&E morphology with VisiumHD expression and expert labels.",
        13,
        NAVY,
    )
    add_metric(slide, Inches(6.25), Inches(1.48), Inches(1.55), Inches(0.88), "448K", "8 um bins", TEAL)
    add_metric(slide, Inches(7.98), Inches(1.48), Inches(1.55), Inches(0.88), "18K", "genes", BLUE)
    add_metric(slide, Inches(9.72), Inches(1.48), Inches(1.55), Inches(0.88), "8 / 3", "labels tested", AMBER)
    add_metric(slide, Inches(11.45), Inches(1.48), Inches(1.15), Inches(0.88), "4", "runs done", GREEN)

    # Flow diagram.
    y = Inches(4.05)
    add_label(slide, Inches(0.78), y, Inches(2.15), Inches(0.68), "H&E crop\n64 or 128 px", BLUE)
    add_label(slide, Inches(3.35), y, Inches(2.15), Inches(0.68), "Image encoder\nResNet / GigaPath", BLUE)
    add_label(slide, Inches(5.92), y, Inches(1.75), Inches(0.68), "128-dim\nembedding", TEAL)
    add_label(slide, Inches(0.78), Inches(5.12), Inches(2.15), Inches(0.68), "Expression\n18,085 genes", AMBER)
    add_label(slide, Inches(3.35), Inches(5.12), Inches(2.15), Inches(0.68), "Expression\nMLP encoder", AMBER)
    add_label(slide, Inches(5.92), Inches(5.12), Inches(1.75), Inches(0.68), "128-dim\nembedding", TEAL)
    add_label(slide, Inches(8.25), Inches(4.55), Inches(2.0), Inches(0.78), "Shared space\nretrieval", GREEN)
    add_label(slide, Inches(10.75), Inches(4.55), Inches(1.75), Inches(0.78), "SAM3 prompt\nfuture", NAVY)
    for x1, y1, x2, y2 in [
        (2.93, 4.39, 3.35, 4.39), (5.50, 4.39, 5.92, 4.39), (7.67, 4.39, 8.25, 4.78),
        (2.93, 5.46, 3.35, 5.46), (5.50, 5.46, 5.92, 5.46), (7.67, 5.46, 8.25, 5.05),
        (10.25, 4.94, 10.75, 4.94),
    ]:
        line = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        line.line.color.rgb = GRAY
        line.line.width = Pt(1.4)


def slide_objectives(prs):
    slide = blank(prs)
    add_title(slide, "Training objective: three signals, not one black box", "The current baseline uses supervised semantics plus a lightweight cross-modal alignment term")
    add_footer(slide)
    add_text(slide, Inches(0.72), Inches(1.48), Inches(11.2), Inches(0.42), "Current completed baseline", 15, NAVY, True)
    add_text(
        slide,
        Inches(0.72),
        Inches(1.95),
        Inches(11.7),
        Inches(0.45),
        "Total loss = CE_image + CE_expression + 0.5 x MSE(normalized image embedding, normalized expression embedding)",
        18,
        BLUE,
        True,
        PP_ALIGN.CENTER,
    )
    cards = [
        ("CE_image", "Can H&E morphology predict the expert tissue label?", BLUE),
        ("CE_expression", "Can gene expression predict the same tissue label?", AMBER),
        ("Alignment", "Are the image and expression embeddings in the same space?", TEAL),
    ]
    for i, (title, body, color) in enumerate(cards):
        x = Inches(0.8 + i * 4.05)
        add_label(slide, x, Inches(3.05), Inches(3.35), Inches(0.62), title, color, RGBColor(239, 246, 255), 15)
        add_text(slide, x, Inches(3.88), Inches(3.35), Inches(0.76), body, 12, NAVY, False, PP_ALIGN.CENTER)
    add_picture(
        slide,
        FIG_DIR / "presentation_3class_total_loss_decomposition.png",
        Inches(2.15),
        Inches(4.82),
        w=Inches(9.1),
        h=Inches(1.95),
    )


def slide_ours_what_it_trains_detailed(prs):
    slide = blank(prs)
    add_title(
        slide,
        "Our training: what the model is actually learning",
        "A two-modal H&E-expression shared-space model before SAM3 integration",
    )
    add_footer(slide)
    add_text(
        slide,
        Inches(0.75),
        Inches(1.35),
        Inches(11.8),
        Inches(0.55),
        "Training sample = one spatial tissue location represented as (H&E crop, VisiumHD expression vector, tissue label)",
        16,
        BLUE,
        True,
        PP_ALIGN.CENTER,
    )
    rows = [
        ["H&E crop", "local morphology around the VisiumHD spot/bin", "image tower learns pathology morphology"],
        ["Expression vector", "18,085-gene log-normalized VisiumHD profile", "expression tower learns molecular / semantic tissue signal"],
        ["Tissue label", "3-class coarse or 8-class fine expert label", "semantic anchor used by both towers through CE losses"],
    ]
    add_table(
        slide,
        Inches(0.65),
        Inches(2.18),
        Inches(12.05),
        Inches(1.55),
        ["Input", "What it is", "What it teaches"],
        rows,
        [Inches(1.65), Inches(4.15), Inches(6.25)],
        8.6,
    )
    goal_rows = [
        ["Main representation goal", "H&E and expression embeddings from biologically similar regions should live in a shared space"],
        ["Classification goal", "image branch and expression branch should both predict the tissue label"],
        ["Alignment goal", "paired H&E-expression embeddings should be closer than mismatched pairs when InfoNCE is used"],
        ["SAM3 goal", "use the learned space to rank regions, exemplars, prompts, or masks before/inside SAM3"],
    ]
    add_table(
        slide,
        Inches(0.9),
        Inches(4.25),
        Inches(11.5),
        Inches(1.8),
        ["Goal", "Meaning"],
        goal_rows,
        [Inches(2.2), Inches(9.3)],
        8.3,
    )
    add_text(
        slide,
        Inches(0.95),
        Inches(6.35),
        Inches(11.35),
        Inches(0.42),
        "Important: this is not SAM3 fine-tuning yet. It is the expression-aware representation that will later feed SAM3 ranking/prompting.",
        12.2,
        NAVY,
        True,
        PP_ALIGN.CENTER,
    )


def slide_ours_architecture_detailed(prs):
    slide = blank(prs)
    add_title(
        slide,
        "Our architecture: two towers, projection heads, shared classifier",
        "This is the two-modal counterpart to Haiku's three-tower shared-space design",
    )
    add_footer(slide)
    add_label(slide, Inches(0.75), Inches(1.55), Inches(1.75), Inches(0.58), "H&E crop\n128 px", BLUE)
    add_label(slide, Inches(2.95), Inches(1.55), Inches(2.15), Inches(0.58), "Image encoder\nGigaPath / ResNet", BLUE)
    add_label(slide, Inches(5.55), Inches(1.55), Inches(1.75), Inches(0.58), "Image\nprojection", TEAL)
    add_label(slide, Inches(0.75), Inches(2.62), Inches(1.75), Inches(0.58), "Expression\n18,085 genes", AMBER)
    add_label(slide, Inches(2.95), Inches(2.62), Inches(2.15), Inches(0.58), "Expression\nMLP encoder", AMBER)
    add_label(slide, Inches(5.55), Inches(2.62), Inches(1.75), Inches(0.58), "Expression\nprojection", TEAL)
    add_label(slide, Inches(8.05), Inches(2.05), Inches(2.0), Inches(0.72), "Shared\nembedding space", NAVY)
    add_label(slide, Inches(10.75), Inches(2.05), Inches(1.65), Inches(0.72), "Shared\nclassifier", GREEN)
    for y in [1.84, 2.91]:
        for x1, x2 in [(2.5, 2.95), (5.1, 5.55), (7.3, 8.05), (10.05, 10.75)]:
            line = slide.shapes.add_connector(1, Inches(x1), Inches(y), Inches(x2), Inches(2.41 if x2 in [8.05, 10.75] else y))
            line.line.color.rgb = GRAY
            line.line.width = Pt(1.2)
    rows = [
        ["Image encoder", "GigaPath tile encoder or ResNet50", "tests pathology foundation model vs generic CNN baseline"],
        ["Expression encoder", "MLP over 18,085 genes", "learns compact expression representation from VisiumHD profile"],
        ["Projection heads", "map both modalities to shared embedding", "the space used by retrieval and alignment losses"],
        ["Shared classifier", "same classifier applied to image/expression embeddings", "forces both modalities to predict the same tissue semantics"],
    ]
    add_table(
        slide,
        Inches(0.65),
        Inches(4.15),
        Inches(12.05),
        Inches(1.72),
        ["Part", "Implementation", "Why it is there"],
        rows,
        [Inches(1.65), Inches(3.45), Inches(6.95)],
        7.7,
    )


def slide_ours_input_construction_detailed(prs):
    slide = blank(prs)
    add_title(
        slide,
        "Our input construction: paired H&E morphology and VisiumHD expression",
        "The training unit is a co-localized image-expression pair with tissue-label supervision",
    )
    add_footer(slide)
    rows = [
        ["H&E source", "Exp1 tissue_hires_image.png", "same histology image underlying VisiumHD spots"],
        ["Image crop", "crop size 128 px for main GigaPath/ResNet ablation", "local morphology around each spatial bin"],
        ["Model input size", "GigaPath resized to 224; ResNet kept at 128", "backbone-specific image input requirement"],
        ["Expression source", "expr_log1p.npz with 18,085 genes", "molecular profile paired with the image crop"],
        ["Normalization", "scaler.npz applied to expression values", "keeps expression features numerically stable"],
        ["Labels", "3-class or 8-class expert tissue labels", "semantic supervision and evaluation target"],
    ]
    add_table(
        slide,
        Inches(0.55),
        Inches(1.45),
        Inches(12.25),
        Inches(2.6),
        ["Input component", "Our setting", "Meaning"],
        rows,
        [Inches(1.75), Inches(4.45), Inches(6.05)],
        7.7,
    )
    add_text(slide, Inches(0.85), Inches(4.6), Inches(3.6), Inches(0.35), "3-class target", 14, NAVY, True)
    add_text(slide, Inches(0.85), Inches(5.0), Inches(4.8), Inches(0.8), "tumor / stroma / immune infiltration\nCoarse biological categories; currently the most stable setting.", 12, NAVY)
    add_text(slide, Inches(6.45), Inches(4.6), Inches(3.6), Inches(0.35), "8-class target", 14, NAVY, True)
    add_text(slide, Inches(6.45), Inches(5.0), Inches(5.1), Inches(0.8), "all expert labels, including rarer or more ambiguous classes\nHarder because labels are finer and less morphologically separable.", 12, NAVY)


def slide_ours_projection_loss_detailed(prs):
    slide = blank(prs)
    add_title(
        slide,
        "Our projection and loss: how the model is optimized",
        "We compare label supervision alone, label supervision plus alignment, and alignment alone",
    )
    add_footer(slide)
    add_text(
        slide,
        Inches(0.75),
        Inches(1.32),
        Inches(11.8),
        Inches(0.5),
        "Total loss variants are built from CE_image, CE_expression, and an image-expression alignment term",
        15,
        BLUE,
        True,
        PP_ALIGN.CENTER,
    )
    rows = [
        ["CE_image", "cross entropy from image embedding -> tissue label", "teaches H&E morphology to predict tissue semantics"],
        ["CE_expression", "cross entropy from expression embedding -> tissue label", "turns expression tower into supervised molecular/semantic anchor"],
        ["MSE alignment", "distance between normalized image/expression embeddings", "baseline pairwise closeness objective"],
        ["InfoNCE alignment", "paired image-expression positives vs in-batch negatives", "CLIP-style alignment; directly optimizes retrieval structure"],
    ]
    add_table(
        slide,
        Inches(0.55),
        Inches(2.05),
        Inches(12.25),
        Inches(1.85),
        ["Loss part", "Definition", "Role"],
        rows,
        [Inches(1.55), Inches(4.45), Inches(6.25)],
        7.6,
    )
    variant_rows = [
        ["Original baseline", "CE_img + CE_expr + 0.5 x MSE", "first working shared-space baseline"],
        ["CE-only", "CE_img + CE_expr", "tests whether shared labels alone create useful retrieval"],
        ["CE + InfoNCE", "CE_img + CE_expr + lambda x symmetric InfoNCE", "tests whether paired alignment adds value beyond labels"],
        ["InfoNCE-only", "symmetric InfoNCE only", "tests whether paired data can align without tissue-label anchor"],
    ]
    add_table(
        slide,
        Inches(0.75),
        Inches(4.45),
        Inches(11.85),
        Inches(1.75),
        ["Experiment", "Loss formula", "Question answered"],
        variant_rows,
        [Inches(1.85), Inches(3.7), Inches(6.3)],
        7.55,
    )
    add_text(slide, Inches(0.9), Inches(6.5), Inches(11.4), Inches(0.32), "InfoNCE uses tau = 0.07; lambda sweep = 0, 0.05, 0.10, 0.20, 0.50.", 11.5, NAVY, True, PP_ALIGN.CENTER)


def slide_ours_training_strategy_detailed(prs):
    slide = blank(prs)
    add_title(
        slide,
        "Our training strategy: what is frozen, what is trained, and what changed later",
        "This is the practical methods table for the runs in the result section",
    )
    add_footer(slide)
    rows = [
        ["Base ablation backbone", "frozen GigaPath / frozen ResNet", "avoid overfitting while testing projection/expression/loss design"],
        ["Trainable in base ablation", "image projection, expression MLP, shared classifier", "these parts learn the cross-modal task"],
        ["Base optimizer", "AdamW, weight decay 1e-4", "standard stable optimizer for neural fine-tuning"],
        ["Base LR", "projection / expression / classifier = 1e-4", "same LR for newly trained parts"],
        ["Regularization", "dropout = 0.3; early stop patience = 2", "reduce overfitting on limited labels"],
        ["Class imbalance", "balanced sampler + sqrt inverse class weights", "rare labels contribute more without exploding weights"],
    ]
    add_table(
        slide,
        Inches(0.55),
        Inches(1.35),
        Inches(12.25),
        Inches(2.55),
        ["Training part", "Our setting", "Why"],
        rows,
        [Inches(2.05), Inches(4.05), Inches(6.15)],
        7.45,
    )
    strategy_rows = [
        ["Frozen constant", "0 trainable blocks", "constant LR", "control"],
        ["Frozen warmup/cosine", "0 trainable blocks", "warmup/cosine", "tests schedule only"],
        ["Last-1 warmup/cosine", "last 1 image block", "backbone LR 1e-5", "controlled pathology feature adaptation"],
        ["Last-2 warmup/cosine", "last 2 image blocks", "backbone LR 1e-5", "closest to Haiku-style image fine-tuning"],
    ]
    add_table(
        slide,
        Inches(0.75),
        Inches(4.45),
        Inches(11.85),
        Inches(1.72),
        ["Training strategy", "Image encoder", "Schedule/LR", "Purpose"],
        strategy_rows,
        [Inches(2.25), Inches(2.25), Inches(2.45), Inches(4.9)],
        7.35,
    )


def slide_haiku_summary(prs):
    slide = blank(prs)
    add_title(
        slide,
        "Haiku paper: align morphology, spatial biology, and clinical semantics",
        "Li et al., arXiv Apr 30 2026: tri-modal contrastive learning for H&E, mIF, and text",
    )
    add_footer(slide)
    add_text(slide, Inches(0.72), Inches(1.42), Inches(4.4), Inches(0.35), "Core idea", 15, NAVY, True)
    add_text(
        slide,
        Inches(0.72),
        Inches(1.86),
        Inches(4.8),
        Inches(1.35),
        "Haiku is not primarily a classifier. It learns one shared embedding space where the same tissue region is close across H&E morphology, mIF spatial proteomics, and clinical/semantic text.",
        13,
        NAVY,
    )
    add_label(slide, Inches(6.0), Inches(1.55), Inches(1.85), Inches(0.62), "H&E patch\nmorphology", BLUE)
    add_label(slide, Inches(8.35), Inches(1.55), Inches(1.85), Inches(0.62), "mIF patch\nproteomics", AMBER)
    add_label(slide, Inches(10.7), Inches(1.55), Inches(1.85), Inches(0.62), "Text\nclinical semantics", GREEN)
    add_label(slide, Inches(7.65), Inches(2.72), Inches(3.2), Inches(0.65), "Shared 512-dim embedding space", TEAL)
    for x in [6.92, 9.27, 11.62]:
        line = slide.shapes.add_connector(1, Inches(x), Inches(2.17), Inches(9.25), Inches(2.72))
        line.line.color.rgb = GRAY
        line.line.width = Pt(1.2)
    rows = [
        ["Training sample", "(H&E patch, mIF patch, text description)", "same spatial location"],
        ["Image pairing", "H&E and mIF patches are registered", "positive pair = same tissue region"],
        ["Text source", "biomarker pattern text + clinical metadata", "structured natural language, not free chat"],
        ["Goal", "retrieval, fusion inference, clinical prediction", "not just image classification"],
    ]
    add_table(
        slide,
        Inches(0.72),
        Inches(3.65),
        Inches(11.8),
        Inches(1.9),
        ["Part", "What Haiku uses", "Why it matters"],
        rows,
        [Inches(1.65), Inches(4.35), Inches(5.8)],
        8.2,
    )
    add_text(
        slide,
        Inches(0.9),
        Inches(6.08),
        Inches(11.2),
        Inches(0.38),
        "One-sentence takeaway: Haiku uses contrastive alignment to make morphology, molecular biology, and clinical semantics mutually searchable.",
        12.5,
        BLUE,
        True,
        PP_ALIGN.CENTER,
    )


def slide_haiku_what_it_does_detailed(prs):
    slide = blank(prs)
    add_title(
        slide,
        "Haiku: what it is actually training",
        "A tri-modal CLIP-style framework over co-registered H&E, mIF, and structured text",
    )
    add_footer(slide)
    add_text(
        slide,
        Inches(0.72),
        Inches(1.38),
        Inches(11.9),
        Inches(0.58),
        "Training sample = one spatial tissue region represented three ways: (H&E patch, mIF patch, text description)",
        17,
        BLUE,
        True,
        PP_ALIGN.CENTER,
    )
    rows = [
        ["H&E patch", "local morphology", "strictly registered to the same spatial location as mIF"],
        ["mIF patch", "spatial proteomics / biomarker image", "multi-channel molecular image, not ordinary RGB"],
        ["Text description", "biomarker pattern + clinical metadata", "generated from patch-level molecular statistics plus patient/slice metadata"],
    ]
    add_table(
        slide,
        Inches(0.65),
        Inches(2.25),
        Inches(12.05),
        Inches(1.55),
        ["Modality", "What it carries", "Where it comes from"],
        rows,
        [Inches(1.55), Inches(3.1), Inches(7.4)],
        8.7,
    )
    goal_rows = [
        ["Alignment goal", "same tissue region across three modalities should be close in embedding space"],
        ["Separation goal", "different regions in the batch should be pushed apart"],
        ["Not the goal", "not just one classifier; classification is only one downstream readout"],
        ["What it enables", "cross-modal retrieval, zero-shot biomarker inference, clinical prediction, counterfactual analysis"],
    ]
    add_table(
        slide,
        Inches(0.95),
        Inches(4.35),
        Inches(11.35),
        Inches(1.7),
        ["Concept", "Meaning"],
        goal_rows,
        [Inches(2.0), Inches(9.35)],
        8.5,
    )
    add_text(
        slide,
        Inches(0.9),
        Inches(6.42),
        Inches(11.4),
        Inches(0.42),
        "This is the central connection to our project: use a shared space first, then use retrieval/fusion to guide biological inference or SAM3 ranking.",
        12.5,
        NAVY,
        True,
        PP_ALIGN.CENTER,
    )


def slide_haiku_architecture_detailed(prs):
    slide = blank(prs)
    add_title(
        slide,
        "Haiku architecture: three towers project into one 512-dim space",
        "Think of it as CLIP extended from two modalities to three modalities",
    )
    add_footer(slide)
    add_label(slide, Inches(0.75), Inches(1.52), Inches(1.8), Inches(0.58), "H&E patch", BLUE)
    add_label(slide, Inches(3.0), Inches(1.52), Inches(1.9), Inches(0.58), "MUSK ViT\nH&E encoder", BLUE)
    add_label(slide, Inches(5.35), Inches(1.52), Inches(1.8), Inches(0.58), "H&E projection", TEAL)
    add_label(slide, Inches(0.75), Inches(2.42), Inches(1.8), Inches(0.58), "mIF patch", AMBER)
    add_label(slide, Inches(3.0), Inches(2.42), Inches(1.9), Inches(0.58), "VirTues\nmIF encoder", AMBER)
    add_label(slide, Inches(5.35), Inches(2.42), Inches(1.8), Inches(0.58), "mIF projection", TEAL)
    add_label(slide, Inches(0.75), Inches(3.32), Inches(1.8), Inches(0.58), "Text", GREEN)
    add_label(slide, Inches(3.0), Inches(3.32), Inches(1.9), Inches(0.58), "BiomedBERT\ntext encoder", GREEN)
    add_label(slide, Inches(5.35), Inches(3.32), Inches(1.8), Inches(0.58), "Text projection", TEAL)
    add_label(slide, Inches(8.05), Inches(2.42), Inches(2.6), Inches(0.78), "Shared 512-dim\nL2-normalized embedding", NAVY)
    for y in [1.81, 2.71, 3.61]:
        for x1, x2 in [(2.55, 3.0), (4.9, 5.35), (7.15, 8.05)]:
            line = slide.shapes.add_connector(1, Inches(x1), Inches(y), Inches(x2), Inches(y if x2 != 8.05 else 2.81))
            line.line.color.rgb = GRAY
            line.line.width = Pt(1.2)
    rows = [
        ["H&E encoder", "pretrained MUSK vision transformer", "pathology morphology foundation model"],
        ["Text encoder", "BiomedBERT", "clinical / biomarker semantic text"],
        ["mIF encoder", "VirTues pretrained on large mIF data", "molecular/protein imaging anchor"],
        ["mIF channel identity", "ESM-3 protein embedding per biomarker channel", "lets model know which protein each channel represents"],
        ["Non-protein channel", "DAPI gets a learned embedding", "handles channels without protein sequence"],
    ]
    add_table(
        slide,
        Inches(0.65),
        Inches(4.45),
        Inches(12.05),
        Inches(1.7),
        ["Part", "Exact encoder / mechanism", "Why it matters"],
        rows,
        [Inches(1.55), Inches(4.05), Inches(6.45)],
        7.25,
    )


def slide_haiku_input_construction(prs):
    slide = blank(prs)
    add_title(
        slide,
        "Haiku input construction: the text is generated from biology, not arbitrary captions",
        "This matters because the text branch carries structured molecular and clinical meaning",
    )
    add_footer(slide)
    rows = [
        ["H&E / mIF patches", "256 x 256 patches; mIF stride floor(0.7P)", "paired local tissue regions"],
        ["Patch filtering", "keep tissue coverage > 90%", "avoid background-dominated positives"],
        ["mIF preprocessing", "per-channel normalize, clip, rescale to [0,1], 8-bit", "make biomarker channels comparable"],
        ["Biomarker text", "mean intensity, slice z-score, percentile", "high/low relative molecular signal"],
        ["Spatial text", "coverage, CV, clustering, GLCM texture", "sparse / uniform / clustered / heterogeneous"],
        ["Clinical text", "organ, disease, stage, survival, treatment response", "patient/slice-level semantic context"],
    ]
    add_table(
        slide,
        Inches(0.65),
        Inches(1.45),
        Inches(12.0),
        Inches(2.9),
        ["Input part", "Construction", "Meaning"],
        rows,
        [Inches(1.75), Inches(4.0), Inches(6.25)],
        7.75,
    )
    add_text(slide, Inches(0.8), Inches(4.85), Inches(3.0), Inches(0.35), "Template logic", 15, NAVY, True)
    add_text(
        slide,
        Inches(0.8),
        Inches(5.3),
        Inches(5.3),
        Inches(0.85),
        "Patch-level biomarker intensity and spatial pattern are translated into structured text, then combined with slice/patient metadata.",
        12.5,
        NAVY,
    )
    add_text(slide, Inches(6.85), Inches(4.85), Inches(3.0), Inches(0.35), "Why we care", 15, NAVY, True)
    add_text(
        slide,
        Inches(6.85),
        Inches(5.3),
        Inches(5.15),
        Inches(0.85),
        "For our project, VisiumHD expression can play the molecular side; later we can add structured region/gene-set text instead of treating labels as the only semantics.",
        12.5,
        NAVY,
    )


def slide_haiku_input_construction_detailed(prs):
    slide = blank(prs)
    add_title(
        slide,
        "Haiku input construction: how the text and mIF signals are built",
        "The text branch is structured biological language generated from quantitative mIF and clinical metadata",
    )
    add_footer(slide)
    image_rows = [
        ["Patch size", "H&E and mIF both 256 x 256", "same scale local region"],
        ["mIF sliding window", "stride = floor(0.7P), about 179", "overlapping tissue coverage"],
        ["Patch filter", "tissue coverage > 90%", "removes background-heavy samples"],
        ["mIF normalization", "per-channel normalize, clip, rescale to [0,1], 8-bit", "standardizes biomarker channels"],
    ]
    add_table(
        slide,
        Inches(0.55),
        Inches(1.35),
        Inches(12.25),
        Inches(1.75),
        ["Image part", "Haiku setting", "Reason"],
        image_rows,
        [Inches(1.6), Inches(4.55), Inches(6.1)],
        7.8,
    )
    text_rows = [
        ["Biomarker intensity", "patch mean intensity", "local protein abundance"],
        ["Relative level", "slice z-score + percentile", "whether marker is high/low within the slide"],
        ["Spatial statistics", "coverage, CV, clustering, GLCM texture", "how signal is distributed spatially"],
        ["Discrete pattern", "sparse / uniform / clustered / heterogeneous", "turns quantitative image statistics into language"],
        ["Clinical metadata", "organ, disease, stage, survival, treatment response", "patient/slice context"],
    ]
    add_table(
        slide,
        Inches(0.55),
        Inches(3.55),
        Inches(12.25),
        Inches(2.35),
        ["Text source", "How it is computed", "What it contributes"],
        text_rows,
        [Inches(1.65), Inches(4.55), Inches(6.05)],
        7.35,
    )
    add_text(
        slide,
        Inches(0.9),
        Inches(6.34),
        Inches(11.4),
        Inches(0.4),
        "Key point for our project: we can later generate analogous structured text from gene sets, tissue labels, or clinical metadata instead of using only class labels.",
        12,
        BLUE,
        True,
        PP_ALIGN.CENTER,
    )


def slide_haiku_projection_loss_detailed(prs):
    slide = blank(prs)
    add_title(
        slide,
        "Haiku projection head and loss: what actually enters contrastive learning",
        "Encoder features are not compared directly; each modality first goes through its own MLP projection head",
    )
    add_footer(slide)
    add_text(
        slide,
        Inches(0.75),
        Inches(1.35),
        Inches(5.6),
        Inches(0.55),
        "Projection head for each modality: g_m(h) = BN(W2 ReLU(W1 h))",
        15,
        BLUE,
        True,
        PP_ALIGN.CENTER,
    )
    projection_rows = [
        ["Architecture", "2-layer MLP", "same structure, separate parameters for each modality"],
        ["Output dimension", "512", "shared embedding dimension"],
        ["Normalization", "L2 normalize after projection", "cosine similarity becomes the comparison metric"],
        ["Compared features", "projected embeddings z, not raw encoder features", "prevents encoder-specific feature scales from dominating"],
    ]
    add_table(
        slide,
        Inches(0.65),
        Inches(2.18),
        Inches(5.95),
        Inches(1.8),
        ["Projection detail", "Value", "Meaning"],
        projection_rows,
        [Inches(1.55), Inches(1.55), Inches(2.85)],
        7.45,
    )
    add_text(
        slide,
        Inches(7.05),
        Inches(1.35),
        Inches(5.0),
        Inches(0.72),
        "InfoNCE: positive = same spatial location; negatives = other batch regions",
        14.5,
        BLUE,
        True,
        PP_ALIGN.CENTER,
    )
    loss_rows = [
        ["Similarity", "cosine similarity", "after L2 normalization"],
        ["Temperature", "tau = 0.07", "controls contrastive sharpness"],
        ["Batch size", "B = 128", "provides in-batch negatives"],
        ["Formula", "exp(sim(pos)/tau) / sum_j exp(sim(neg_j)/tau)", "correct pair should rank above wrong pairs"],
    ]
    add_table(
        slide,
        Inches(6.95),
        Inches(2.18),
        Inches(5.55),
        Inches(1.8),
        ["Loss detail", "Value", "Meaning"],
        loss_rows,
        [Inches(1.25), Inches(1.75), Inches(2.55)],
        7.25,
    )
    pair_rows = [
        ["L_HE <-> mIF", "morphology to spatial protein pattern"],
        ["L_mIF <-> TXT", "protein pattern to structured clinical/biomarker semantics"],
        ["L_HE <-> TXT", "morphology to clinical/semantic context"],
    ]
    add_table(
        slide,
        Inches(2.3),
        Inches(4.75),
        Inches(8.7),
        Inches(1.3),
        ["Symmetric loss pair", "What it aligns"],
        pair_rows,
        [Inches(2.2), Inches(6.5)],
        8.2,
    )
    add_text(
        slide,
        Inches(0.85),
        Inches(6.33),
        Inches(11.6),
        Inches(0.4),
        "This is why we replaced simple MSE alignment with CLIP-style symmetric InfoNCE in our ablation.",
        12.2,
        NAVY,
        True,
        PP_ALIGN.CENTER,
    )


def slide_haiku_training_strategy(prs):
    slide = blank(prs)
    add_title(
        slide,
        "Literature anchor: Haiku trains alignment around a frozen molecular encoder",
        "This is the closest paper to our H&E-expression-text setting and motivates the ablation design",
    )
    add_footer(slide)
    add_text(slide, Inches(0.72), Inches(1.42), Inches(4.0), Inches(0.35), "Haiku tri-modal alignment", 15, NAVY, True)
    add_label(slide, Inches(0.85), Inches(2.0), Inches(2.25), Inches(0.68), "mIF encoder\nfrozen anchor", AMBER)
    add_label(slide, Inches(3.75), Inches(1.55), Inches(2.25), Inches(0.68), "H&E encoder\nfine-tune last 2 blocks", BLUE)
    add_label(slide, Inches(3.75), Inches(2.45), Inches(2.25), Inches(0.68), "Text encoder\nfine-tune last 2 blocks", GREEN)
    add_label(slide, Inches(6.65), Inches(2.0), Inches(2.2), Inches(0.68), "Projection heads\nfully trained", TEAL)
    add_label(slide, Inches(9.45), Inches(2.0), Inches(2.25), Inches(0.68), "Shared 512-dim\nnormalized space", NAVY)
    for x1, y1, x2, y2 in [
        (3.10, 2.34, 3.75, 1.89), (3.10, 2.34, 3.75, 2.79),
        (6.00, 1.89, 6.65, 2.34), (6.00, 2.79, 6.65, 2.34),
        (8.85, 2.34, 9.45, 2.34),
    ]:
        line = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        line.line.color.rgb = GRAY
        line.line.width = Pt(1.4)

    rows = [
        ["Frozen module", "mIF encoder", "acts as a pretrained molecular anchor"],
        ["Fine-tuned modules", "H&E + text encoders", "only last two transformer blocks"],
        ["Fully trained", "projection heads", "map modalities into shared 512-dim space"],
        ["Optimizer / schedule", "AdamW", "5,000-step warmup then cosine annealing"],
        ["Learning rates", "1e-5 / 2e-5 / 1e-4", "H&E / text / projection heads"],
        ["Training length", "25 epochs", "alignment stage, not end-to-end early fusion"],
    ]
    add_table(
        slide,
        Inches(0.72),
        Inches(3.65),
        Inches(6.2),
        Inches(2.55),
        ["Strategy", "Haiku setting", "Interpretation"],
        rows,
        [Inches(1.45), Inches(1.75), Inches(3.0)],
        7.7,
    )
    add_text(slide, Inches(7.35), Inches(3.7), Inches(4.9), Inches(0.35), "How we translate this", 15, NAVY, True)
    add_text(
        slide,
        Inches(7.35),
        Inches(4.18),
        Inches(4.95),
        Inches(1.6),
        "Our expression tower is not a pretrained mIF encoder, so we first use CE_expression as a supervised molecular/semantic anchor. Then we test whether adding CLIP-style InfoNCE improves the shared H&E-expression space beyond label supervision alone.",
        12.5,
        NAVY,
    )
    add_text(
        slide,
        Inches(7.35),
        Inches(5.9),
        Inches(4.95),
        Inches(0.42),
        "This is why the ablation compares CE-only, CE+InfoNCE, and InfoNCE-only.",
        12.5,
        BLUE,
        True,
    )


def slide_haiku_training_fusion_tasks_detailed(prs):
    slide = blank(prs)
    add_title(
        slide,
        "Haiku training, fusion, and final tasks: the methods details we borrow",
        "The paper's main lesson is controlled fine-tuning plus retrieval-time fusion, not early-fusion classification",
    )
    add_footer(slide)
    train_rows = [
        ["Frozen", "mIF encoder", "pretrained molecular anchor"],
        ["Fine-tuned", "last 2 blocks of H&E encoder", "small adaptation of MUSK morphology features"],
        ["Fine-tuned", "last 2 blocks of text encoder", "small adaptation of BiomedBERT semantics"],
        ["Fully trained", "all projection heads", "main learnable bridge into shared 512-dim space"],
        ["Learning rates", "H&E 1e-5; text 2e-5; projection 1e-4", "small LR for foundation encoders, larger LR for new heads"],
        ["Optimizer/schedule", "AdamW; 5000 warmup steps; cosine annealing; 25 epochs", "stable contrastive alignment"],
    ]
    add_table(
        slide,
        Inches(0.5),
        Inches(1.32),
        Inches(12.35),
        Inches(2.52),
        ["Training part", "Haiku choice", "Interpretation"],
        train_rows,
        [Inches(1.65), Inches(3.55), Inches(7.15)],
        7.15,
    )
    fusion_rows = [
        ["Training-time fusion", "none / shared-space alignment", "modalities are not concatenated as early input"],
        ["Biomarker inference", "0.8 x H&E + 0.2 x text", "mostly morphology, some clinical/semantic conditioning"],
        ["Counterfactual retrieval", "0.6 x H&E + 0.4 x text", "increase text effect to test conditional shifts"],
        ["Our equivalent", "image-expression alpha sweep", "rank H&E regions / candidate SAM3 masks using fused evidence"],
    ]
    add_table(
        slide,
        Inches(0.55),
        Inches(4.28),
        Inches(6.0),
        Inches(1.72),
        ["Fusion setting", "Weighting", "Meaning"],
        fusion_rows,
        [Inches(1.8), Inches(1.65), Inches(2.55)],
        7.1,
    )
    task_rows = [
        ["Cross-modal retrieval", "H&E <-> mIF; text -> mIF"],
        ["Downstream prediction", "organ, tissue type, T/N stage, grade, survival, treatment response"],
        ["Zero-shot biomarker inference", "retrieve mIF neighbors and aggregate marker intensity"],
        ["Counterfactual discovery", "fix H&E, edit text metadata, observe retrieved molecular neighbors"],
    ]
    add_table(
        slide,
        Inches(6.85),
        Inches(4.28),
        Inches(5.6),
        Inches(1.72),
        ["Final use", "What it tests"],
        task_rows,
        [Inches(1.8), Inches(3.8)],
        7.1,
    )


def slide_haiku_loss_and_tasks(prs):
    slide = blank(prs)
    add_title(
        slide,
        "Haiku objective: three pairwise contrastive losses create one searchable space",
        "Positive pair = same spatial location; negatives = other batch locations; tau = 0.07; batch size = 128",
    )
    add_footer(slide)
    add_text(
        slide,
        Inches(0.72),
        Inches(1.45),
        Inches(11.8),
        Inches(0.5),
        "For modality pair (a,b):  L_a->b = -log exp(sim(z_i^a,z_i^b)/tau) / sum_j exp(sim(z_i^a,z_j^b)/tau)",
        15,
        BLUE,
        True,
        PP_ALIGN.CENTER,
    )
    rows = [
        ["H&E -> mIF", "align morphology to spatial proteomics", "H&E can retrieve molecular states"],
        ["mIF -> Text", "align protein marker pattern to generated semantics", "text can query molecular atlas"],
        ["H&E -> Text", "align morphology to clinical/semantic description", "H&E and metadata become comparable"],
        ["Symmetric terms", "reverse directions included implicitly", "retrieval works both ways"],
    ]
    add_table(
        slide,
        Inches(0.65),
        Inches(2.35),
        Inches(5.95),
        Inches(2.0),
        ["Loss pair", "What it aligns", "Consequence"],
        rows,
        [Inches(1.25), Inches(2.55), Inches(2.15)],
        7.8,
    )
    task_rows = [
        ["Cross-modal retrieval", "H&E <-> mIF, text -> mIF", "tests whether modalities are truly aligned"],
        ["Downstream prediction", "organ, tissue type, stage, grade, survival, response", "tests representation quality"],
        ["Zero-shot biomarker inference", "retrieve mIF neighbors for an H&E/text query", "estimate biomarker profile without direct regression"],
        ["Counterfactual discovery", "change text metadata while fixing H&E", "hypothesis generation, not causal proof"],
    ]
    add_table(
        slide,
        Inches(6.95),
        Inches(2.35),
        Inches(5.55),
        Inches(2.65),
        ["Final task", "How it is used", "Point"],
        task_rows,
        [Inches(1.55), Inches(2.2), Inches(1.8)],
        7.35,
    )
    add_text(
        slide,
        Inches(0.9),
        Inches(5.65),
        Inches(11.2),
        Inches(0.62),
        "The important design lesson for us: evaluate the shared space with retrieval and fusion, not only validation accuracy.",
        13,
        NAVY,
        True,
        PP_ALIGN.CENTER,
    )


def slide_haiku_fusion_strategy(prs):
    slide = blank(prs)
    add_title(
        slide,
        "Haiku fusion is retrieval-time weighting, not early fusion",
        "Training aligns modalities into one space; inference mixes scores or normalized embeddings",
    )
    add_footer(slide)
    add_text(slide, Inches(0.72), Inches(1.45), Inches(4.5), Inches(0.35), "Training-time idea", 15, NAVY, True)
    add_label(slide, Inches(0.85), Inches(2.05), Inches(1.75), Inches(0.65), "H&E", BLUE)
    add_label(slide, Inches(0.85), Inches(3.0), Inches(1.75), Inches(0.65), "Text metadata", GREEN)
    add_label(slide, Inches(3.2), Inches(2.52), Inches(1.95), Inches(0.65), "Shared space", TEAL)
    add_label(slide, Inches(5.75), Inches(2.52), Inches(1.75), Inches(0.65), "mIF gallery", AMBER)
    for x1, y1, x2, y2 in [(2.60, 2.38, 3.20, 2.85), (2.60, 3.33, 3.20, 2.85), (5.15, 2.85, 5.75, 2.85)]:
        line = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        line.line.color.rgb = GRAY
        line.line.width = Pt(1.4)

    add_text(
        slide,
        Inches(0.72),
        Inches(4.2),
        Inches(5.9),
        Inches(0.85),
        "No early-fusion network is trained. Each modality gets an embedding, and contrastive losses make matched samples closer than mismatched samples.",
        12.5,
        NAVY,
    )
    add_text(slide, Inches(7.25), Inches(1.45), Inches(4.5), Inches(0.35), "Inference-time fusion", 15, NAVY, True)
    add_text(
        slide,
        Inches(7.15),
        Inches(2.05),
        Inches(5.2),
        Inches(0.55),
        "score_fusion = alpha x sim(H&E, mIF) + (1-alpha) x sim(Text, mIF)",
        15,
        BLUE,
        True,
        PP_ALIGN.CENTER,
    )
    rows = [
        ["Biomarker inference", "alpha = 0.8", "H&E 80%, text 20%"],
        ["Counterfactual retrieval", "alpha = 0.6", "H&E 60%, text 40%"],
        ["Our equivalent", "alpha sweep", "image-expression fusion for retrieval/proposal ranking"],
    ]
    add_table(
        slide,
        Inches(7.25),
        Inches(3.05),
        Inches(4.8),
        Inches(1.45),
        ["Use case", "Weight", "Meaning"],
        rows,
        [Inches(1.7), Inches(1.15), Inches(1.95)],
        8.2,
    )
    add_text(
        slide,
        Inches(7.25),
        Inches(5.0),
        Inches(4.9),
        Inches(0.95),
        "For SAM3, this suggests a practical next step: use expression/text-conditioned retrieval to rank H&E regions or candidate masks, instead of forcing SAM3 to directly learn gene expression.",
        12.5,
        NAVY,
    )


def slide_our_training_strategy(prs):
    slide = blank(prs)
    add_title(
        slide,
        "Our exact training strategy: frozen image backbone, train the alignment heads",
        "This is the Methods-level detail for the completed GigaPath/ResNet ablations",
    )
    add_footer(slide)
    add_text(slide, Inches(0.72), Inches(1.42), Inches(4.8), Inches(0.35), "What is trained in our current ablation", 15, NAVY, True)
    add_label(slide, Inches(0.85), Inches(2.0), Inches(2.05), Inches(0.62), "H&E backbone\nfrozen", BLUE)
    add_label(slide, Inches(3.35), Inches(2.0), Inches(2.1), Inches(0.62), "Image projection\ntrained", TEAL)
    add_label(slide, Inches(5.9), Inches(2.0), Inches(2.0), Inches(0.62), "Expression MLP\ntrained", AMBER)
    add_label(slide, Inches(8.35), Inches(2.0), Inches(1.9), Inches(0.62), "Shared classifier\ntrained", GREEN)
    add_label(slide, Inches(10.65), Inches(2.0), Inches(1.65), Inches(0.62), "128-dim\nspace", NAVY)
    for x1, x2 in [(2.90, 3.35), (5.45, 5.90), (7.90, 8.35), (10.25, 10.65)]:
        line = slide.shapes.add_connector(1, Inches(x1), Inches(2.31), Inches(x2), Inches(2.31))
        line.line.color.rgb = GRAY
        line.line.width = Pt(1.4)

    rows = [
        ["Image backbone", "GigaPath or ResNet50", "Frozen for all 8 epochs in the ablation"],
        ["Image projection", "feature_dim -> 256 -> embed_dim", "Trained with AdamW, lr 1e-4"],
        ["Expression encoder", "18,085 genes -> 256 -> 256 -> embed_dim", "Trained with AdamW, lr 1e-4, dropout 0.3"],
        ["Classifier", "shared linear head on image/expression embeddings", "Trained; CE_image and CE_expr both backprop through it"],
        ["Backbone LR", "configured as 1e-5 after unfreeze", "not used in ablation because freeze_epochs = total_epochs"],
        ["Schedule", "constant LR, AdamW, weight decay 1e-4", "no warmup/cosine yet; this differs from Haiku"],
    ]
    add_table(
        slide,
        Inches(0.65),
        Inches(3.12),
        Inches(12.05),
        Inches(2.58),
        ["Component", "Implementation", "Training status"],
        rows,
        [Inches(1.8), Inches(3.65), Inches(6.6)],
        7.65,
    )
    add_text(
        slide,
        Inches(0.9),
        Inches(6.05),
        Inches(11.4),
        Inches(0.45),
        "Why freeze the image backbone first: GigaPath/MUSK-like pathology encoders are already strong; with limited VisiumHD labels, training the projection + expression side is safer and faster before doing selective fine-tuning.",
        12,
        NAVY,
        False,
        PP_ALIGN.CENTER,
    )


def slide_our_pipeline_summary(prs):
    slide = blank(prs)
    add_title(
        slide,
        "Our project: a two-modal version for H&E and VisiumHD expression",
        "Current goal: learn an expression-aware H&E representation that can later guide SAM3 masks/prompts",
    )
    add_footer(slide)
    add_label(slide, Inches(0.8), Inches(1.72), Inches(1.95), Inches(0.62), "H&E crop\n128 px", BLUE)
    add_label(slide, Inches(3.25), Inches(1.72), Inches(2.1), Inches(0.62), "Image encoder\nGigaPath / ResNet", BLUE)
    add_label(slide, Inches(5.85), Inches(1.72), Inches(1.85), Inches(0.62), "Projection\n128-dim", TEAL)
    add_label(slide, Inches(0.8), Inches(3.0), Inches(1.95), Inches(0.62), "Expression\n18,085 genes", AMBER)
    add_label(slide, Inches(3.25), Inches(3.0), Inches(2.1), Inches(0.62), "Expression MLP\ntrained", AMBER)
    add_label(slide, Inches(5.85), Inches(3.0), Inches(1.85), Inches(0.62), "Projection\n128-dim", TEAL)
    add_label(slide, Inches(8.35), Inches(2.35), Inches(1.95), Inches(0.72), "Shared space\nretrieval", GREEN)
    add_label(slide, Inches(10.85), Inches(2.35), Inches(1.55), Inches(0.72), "SAM3\nnext", NAVY)
    for x1, y1, x2, y2 in [
        (2.75, 2.03, 3.25, 2.03), (5.35, 2.03, 5.85, 2.03), (7.70, 2.03, 8.35, 2.64),
        (2.75, 3.31, 3.25, 3.31), (5.35, 3.31, 5.85, 3.31), (7.70, 3.31, 8.35, 2.78),
        (10.30, 2.72, 10.85, 2.72),
    ]:
        line = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        line.line.color.rgb = GRAY
        line.line.width = Pt(1.4)
    rows = [
        ["Molecular modality", "VisiumHD gene expression", "analogous role to Haiku's mIF, but not a pretrained protein image encoder"],
        ["Semantic anchor", "expert tissue labels: 3-class or 8-class", "CE_expr makes expression branch biologically/semantically useful"],
        ["Alignment target", "H&E embedding <-> expression embedding", "retrieval tests whether same tissue location is close"],
        ["Downstream use", "rank regions/masks before SAM3 integration", "do not force SAM3 to directly regress expression yet"],
    ]
    add_table(
        slide,
        Inches(0.8),
        Inches(4.35),
        Inches(11.7),
        Inches(1.75),
        ["Design choice", "Our implementation", "Why"],
        rows,
        [Inches(1.8), Inches(3.3), Inches(6.6)],
        7.9,
    )


def slide_our_experiment_settings(prs):
    slide = blank(prs)
    add_title(
        slide,
        "Our experiment grid: same training recipe, different supervision signals",
        "The ablation asks whether retrieval comes from shared labels, cross-modal alignment, or both",
    )
    add_footer(slide)
    rows = [
        ["3-class", "tumor / stroma / immune infiltration", "coarse labels; expected more stable"],
        ["8-class", "all expert labels", "fine-grained labels; tests rare/ambiguous categories"],
        ["Crop/input", "GigaPath: 128 -> 224; ResNet: 128 -> 128", "same H&E crop, backbone-specific input resize"],
        ["Batch", "GigaPath 32; ResNet 128", "GigaPath is heavier, ResNet allows larger batch"],
        ["Epochs", "8 total, early stop patience 2", "fast controlled ablation, best checkpoint by val macro-F1"],
        ["Sampling/CE", "balanced sampler + sqrt inverse class weights", "reduces dominance of common classes"],
    ]
    add_table(
        slide,
        Inches(0.65),
        Inches(1.45),
        Inches(12.0),
        Inches(2.55),
        ["Setting", "Value", "Reason"],
        rows,
        [Inches(1.5), Inches(3.65), Inches(6.85)],
        7.8,
    )
    add_text(slide, Inches(0.85), Inches(4.45), Inches(3.35), Inches(0.35), "Loss variants", 15, NAVY, True)
    loss_rows = [
        ["CE-only", "CE_img + CE_expr", "Do the two towers only learn label semantics separately?"],
        ["CE + InfoNCE", "CE_img + CE_expr + lambda x symmetric InfoNCE", "Does paired H&E-expression alignment improve the shared space?"],
        ["InfoNCE-only", "symmetric InfoNCE", "Can paired data align without tissue labels?"],
    ]
    add_table(
        slide,
        Inches(0.75),
        Inches(4.95),
        Inches(6.1),
        Inches(1.32),
        ["Variant", "Loss", "Question"],
        loss_rows,
        [Inches(1.25), Inches(2.1), Inches(2.75)],
        7.55,
    )
    add_text(slide, Inches(7.25), Inches(4.45), Inches(4.7), Inches(0.35), "Lambda sweep", 15, NAVY, True)
    for i, lam in enumerate(["0", "0.05", "0.10", "0.20", "0.50", "only"]):
        color = GRAY if lam == "0" else TEAL if lam != "only" else AMBER
        add_label(slide, Inches(7.25 + i * 0.78), Inches(4.95), Inches(0.62), Inches(0.45), lam, color, LIGHT, 9.5)
    add_text(
        slide,
        Inches(7.25),
        Inches(5.65),
        Inches(4.85),
        Inches(0.75),
        "Presentation interpretation: if CE-only is strong but InfoNCE-only fails, labels are the semantic anchor; if CE+InfoNCE improves retrieval, contrastive alignment adds value beyond classification.",
        11.5,
        NAVY,
    )


def slide_training_variants_map(prs):
    slide = blank(prs)
    add_title(
        slide,
        "What we changed across experiments",
        "The goal is to isolate backbone choice, label granularity, alignment loss, and fine-tuning strategy",
    )
    add_footer(slide)
    groups = [
        ("Backbone", BLUE, ["ResNet50", "GigaPath", "MUSK pending"]),
        ("Label target", GREEN, ["3-class coarse", "8-class fine"]),
        ("Loss design", AMBER, ["CE-only", "CE + InfoNCE", "InfoNCE-only"]),
        ("Training strategy", TEAL, ["Frozen", "Last-1 block", "Last-2 blocks"]),
    ]
    for i, (title, color, items) in enumerate(groups):
        x = Inches(0.65 + i * 3.15)
        add_label(slide, x, Inches(1.55), Inches(2.55), Inches(0.55), title, color, LIGHT, 14)
        for j, item in enumerate(items):
            add_label(slide, x + Inches(0.18), Inches(2.35 + j * 0.78), Inches(2.18), Inches(0.48), item, color, WHITE, 11, False)
    add_text(slide, Inches(0.85), Inches(5.25), Inches(11.6), Inches(0.42), "Read the experiment grid from left to right: choose a backbone, choose label granularity, choose the loss, then choose whether to fine-tune the image encoder.", 13, NAVY, True, PP_ALIGN.CENTER)
    add_text(slide, Inches(1.05), Inches(5.98), Inches(11.2), Inches(0.5), "This is why the result section separates supervised metrics from retrieval/fusion metrics: classification tells us whether each tower learned tissue semantics; retrieval tells us whether the two towers share a useful space.", 12.2, NAVY, False, PP_ALIGN.CENTER)


def slide_our_vs_haiku(prs):
    slide = blank(prs)
    add_title(
        slide,
        "How our current plan maps to Haiku",
        "We borrow the shared-space/retrieval idea, but adapt it to expression + SAM3 rather than mIF + clinical text",
    )
    add_footer(slide)
    rows = [
        ["Molecular anchor", "Frozen pretrained mIF encoder", "Expression MLP trained with CE_expr", "Need stronger molecular pretraining/gene-set text later"],
        ["Image encoder", "MUSK, fine-tune last two blocks", "GigaPath/ResNet frozen in ablation; MUSK smoke submitted", "MUSK is worth testing because Haiku used it"],
        ["Text modality", "BiomedBERT from structured biomarker + clinical text", "not in current training", "next: structured gene-set/tissue-region text"],
        ["Projection", "2-layer MLP to 512-dim", "MLP to 128-dim; embed512 GigaPath job running", "test whether 512 helps retrieval"],
        ["Loss", "pairwise symmetric InfoNCE across 3 modality pairs", "CE_img + CE_expr, with MSE/InfoNCE ablation", "CE is needed because expression side lacks pretrained anchor"],
        ["Fusion", "weighted H&E/text query to mIF gallery", "retrieval/fusion alpha sweep for image-expression", "use for SAM3 proposal/mask ranking"],
    ]
    add_table(
        slide,
        Inches(0.45),
        Inches(1.35),
        Inches(12.45),
        Inches(3.35),
        ["Aspect", "Haiku", "Our current setup", "Decision"],
        rows,
        [Inches(1.35), Inches(3.0), Inches(3.55), Inches(4.55)],
        6.9,
    )
    add_text(
        slide,
        Inches(0.8),
        Inches(5.25),
        Inches(11.6),
        Inches(0.95),
        "Message for group meeting: I am not claiming we already replicated Haiku. I am using Haiku as the design template, then testing which parts transfer to H&E + VisiumHD expression before connecting the representation to SAM3.",
        12.5,
        NAVY,
        True,
        PP_ALIGN.CENTER,
    )


def slide_completed_runs(prs):
    slide = blank(prs)
    add_title(slide, "Completed experiments: ResNet baseline and GigaPath variants", "These are the runs already completed before the new loss ablation")
    add_footer(slide)
    rows = [
        ["ResNet contrastive", "11016975", "InfoNCE only", "8", "64/64", "simple cross-modal baseline"],
        ["GigaPath crop64", "11016957", "CE + CE + MSE", "8", "64/224", "pathology image features, local crop"],
        ["GigaPath crop128", "11016973", "CE + CE + MSE", "8", "128/224", "larger context for fine labels"],
        ["GigaPath crop128", "11016974", "CE + CE + MSE", "3", "128/224", "coarse tumor/stroma/immune target"],
    ]
    add_table(
        slide,
        Inches(0.55),
        Inches(1.5),
        Inches(12.2),
        Inches(2.3),
        ["Run", "Job", "Loss", "Classes", "Crop/Input", "Why it was run"],
        rows,
        [Inches(2.1), Inches(1.05), Inches(1.65), Inches(0.8), Inches(1.25), Inches(5.35)],
        8.8,
    )
    add_text(slide, Inches(0.72), Inches(4.22), Inches(3.9), Inches(0.35), "Best supervised result", 14, NAVY, True)
    add_metric(slide, Inches(0.72), Inches(4.65), Inches(1.55), Inches(0.85), "0.765", "last-2 expr acc", BLUE)
    add_metric(slide, Inches(2.45), Inches(4.65), Inches(1.55), Inches(0.85), "0.694", "last-2 macro-F1", TEAL)
    add_metric(slide, Inches(4.18), Inches(4.65), Inches(1.55), Inches(0.85), "0.745", "CE-only R@1", GREEN)
    add_text(
        slide,
        Inches(6.2),
        Inches(4.45),
        Inches(5.8),
        Inches(1.05),
        "Takeaway: coarse 3-class labels are still more stable than 8-class labels, but Haiku-inspired last-2-block fine-tuning improved both 3-class supervised metrics and 8-class image recognition. Retrieval now shows the fine-tuned models help fused/probe readouts more than raw R@1.",
        13,
        NAVY,
    )


def slide_key_training_result_tables(prs):
    slide = blank(prs)
    add_title(
        slide,
        "Key training result tables: what finished and what it means",
        "These are the numbers to cite verbally before showing retrieval",
    )
    add_footer(slide)
    baseline_rows = [
        ["GigaPath 8c crop64", "11016957", "0.5948", "0.4368", "0.6344", "8-class baseline, smaller crop"],
        ["GigaPath 8c crop128", "11016973", "0.6100", "0.4366", "0.5328", "larger crop did not solve fine labels"],
        ["GigaPath 3c crop128", "11016974", "0.7544", "0.6724", "0.6800", "coarse labels much more stable"],
        ["GigaPath 8c last2 warmcos", "11040134", "0.6316", "0.4512", "0.7084", "fine-tuning improves image branch"],
    ]
    add_table(
        slide,
        Inches(0.45),
        Inches(1.28),
        Inches(12.45),
        Inches(1.85),
        ["Run", "Job", "Expr acc", "Expr macro-F1", "Image acc", "Interpretation"],
        baseline_rows,
        [Inches(2.35), Inches(0.9), Inches(0.85), Inches(1.0), Inches(0.9), Inches(6.45)],
        7.15,
    )
    strategy_rows = [
        ["Frozen constant", "0", "constant", "0.7568", "0.6799", "0.6912"],
        ["Frozen warm/cos", "0", "warmup/cosine", "0.7532", "0.6796", "0.6908"],
        ["Last-1 warm/cos", "1", "warmup/cosine", "0.7652", "0.6929", "0.7920"],
        ["Last-2 warm/cos", "2", "warmup/cosine", "0.7652", "0.6945", "0.8312"],
    ]
    add_text(slide, Inches(0.65), Inches(3.55), Inches(5.5), Inches(0.3), "Haiku-inspired 3-class training strategy", 14, NAVY, True)
    add_table(
        slide,
        Inches(0.45),
        Inches(3.95),
        Inches(6.15),
        Inches(1.75),
        ["Strategy", "Blocks", "LR schedule", "Expr acc", "Macro-F1", "Image acc"],
        strategy_rows,
        [Inches(1.55), Inches(0.65), Inches(1.25), Inches(0.82), Inches(0.9), Inches(0.98)],
        6.9,
    )
    retrieval_rows = [
        ["CE-only", "0.745", "0.811", "0.848", "0.856"],
        ["InfoNCE .10", "0.722", "0.858", "0.900", "0.858"],
        ["Last-1 ft", "0.741", "0.833", "0.870", "0.874"],
        ["Last-2 ft", "0.728", "0.837", "0.865", "0.880"],
    ]
    add_text(slide, Inches(6.95), Inches(3.55), Inches(5.5), Inches(0.3), "3-class retrieval/fusion readout", 14, NAVY, True)
    add_table(
        slide,
        Inches(6.85),
        Inches(3.95),
        Inches(5.65),
        Inches(1.75),
        ["Run", "R@1", "R@5", "R@10", "Fused R@1"],
        retrieval_rows,
        [Inches(1.6), Inches(0.78), Inches(0.78), Inches(0.78), Inches(1.15)],
        7.2,
    )
    add_text(
        slide,
        Inches(0.85),
        Inches(6.18),
        Inches(11.6),
        Inches(0.44),
        "Main interpretation: last-2 fine-tuning wins supervised image accuracy and fused/probe retrieval, while CE-only / last-1 remain strong raw nearest-neighbor baselines.",
        12.2,
        BLUE,
        True,
        PP_ALIGN.CENTER,
    )


def slide_training_curves(prs):
    slide = blank(prs)
    add_title(slide, "Training behavior: 3-class converges; 8-class generalizes poorly", "The 8-class model learns training labels, but validation expression CE stays high")
    add_footer(slide)
    add_picture(slide, FIG_DIR / "presentation_3class_vs_8class_metrics.png", Inches(0.72), Inches(1.45), w=Inches(5.8))
    add_picture(slide, FIG_DIR / "presentation_expression_ce_generalization_gap.png", Inches(6.8), Inches(1.45), w=Inches(5.8))
    add_text(
        slide,
        Inches(1.0),
        Inches(6.15),
        Inches(11.5),
        Inches(0.45),
        "Interpretation: fine 8-class labels likely mix morphology ambiguity, rare classes, and label noise. This is why the next experiment tests whether alignment helps, hurts, or is dominated by classification CE.",
        12.5,
        NAVY,
        False,
        PP_ALIGN.CENTER,
    )


def slide_retrieval_results(prs):
    slide = blank(prs)
    add_title(slide, "Retrieval results: fine-tuning helps fused/probe readouts", "Expression-to-image retrieval checks whether expression embeddings retrieve matching H&E regions")
    add_footer(slide)
    rows = [
        ["GigaPath 3c CE-only", "0.745", "0.811", "0.848", "0.856"],
        ["GigaPath 3c InfoNCE .10", "0.722", "0.858", "0.900", "0.858"],
        ["GigaPath 3c last-1 ft", "0.741", "0.833", "0.870", "0.874"],
        ["GigaPath 3c last-2 ft", "0.728", "0.837", "0.865", "0.880"],
        ["GigaPath 3c InfoNCE .50", "0.721", "0.852", "0.888", "0.861"],
        ["ResNet 3c CE-only", "0.720", "0.868", "0.902", "0.759"],
    ]
    add_table(
        slide,
        Inches(0.65),
        Inches(1.42),
        Inches(5.65),
        Inches(2.78),
        ["Run", "R@1", "R@5", "R@10", "Best fused R@1"],
        rows,
        [Inches(2.3), Inches(0.7), Inches(0.7), Inches(0.75), Inches(1.05)],
        8.5,
    )
    add_picture(slide, FIG_DIR / "retrieval_overall_recall.png", Inches(6.55), Inches(1.25), w=Inches(5.9))
    add_text(slide, Inches(0.9), Inches(4.55), Inches(5.2), Inches(1.2), "Current retrieval readout: CE-only still has the strongest raw 3-class R@1. Last-1 fine-tuning nearly matches it, while last-2 gives the best fused R@1 and concatenated probe, so the shared representation improved even though raw nearest-neighbor R@1 did not.", 12.2, NAVY)
    add_picture(slide, FIG_DIR / "gigapath_3class_per_label_recall1.png", Inches(6.95), Inches(4.25), w=Inches(5.1))


def slide_ablation_plan(prs):
    slide = blank(prs)
    add_title(slide, "New ablation: test what each loss term contributes", "This turns the loss design from an assumption into an experiment")
    add_footer(slide)
    rows = [
        ["CE-only", "CE_img + CE_expr", "completed", "best raw 3-class R@1 so far"],
        ["CE + InfoNCE", "CE_img + CE_expr + lambda x symmetric InfoNCE", "completed", "better top-k/fused retrieval in some runs"],
        ["InfoNCE-only", "symmetric InfoNCE only", "completed", "weak classification, useful diagnostic"],
        ["Haiku strategy", "last-1/2 blocks + warmup/cosine", "completed + retrieval done", "fine-tuning improves fused/probe readouts"],
    ]
    add_table(
        slide,
        Inches(0.65),
        Inches(1.45),
        Inches(12.1),
        Inches(2.45),
        ["Experiment", "Loss", "Status", "Question answered"],
        rows,
        [Inches(1.75), Inches(4.1), Inches(1.2), Inches(5.05)],
        8.7,
    )
    add_text(slide, Inches(0.85), Inches(4.35), Inches(3.6), Inches(0.35), "InfoNCE lambda sweep", 15, NAVY, True)
    for i, lam in enumerate(["0", "0.05", "0.10", "0.20", "0.50"]):
        color = GRAY if lam == "0" else TEAL
        add_label(slide, Inches(0.9 + i * 1.18), Inches(4.9), Inches(0.88), Inches(0.55), lam, color, LIGHT, 13)
    add_text(
        slide,
        Inches(6.2),
        Inches(4.35),
        Inches(5.8),
        Inches(1.2),
        "Overnight readout: training-strategy retrieval completed, the 8-class last-2-block warmup/cosine run completed, and the former OOM retrieval tasks were rerun successfully with chunked retrieval. The validation set is now complete and archived.",
        13,
        NAVY,
    )


def slide_current_hpc_status(prs):
    slide = blank(prs)
    add_title(slide, "Current HPC status: validation set is complete", "Stable CSV/JSON snapshots are written under the project results directory")
    add_footer(slide)
    add_metric(slide, Inches(0.8), Inches(1.55), Inches(1.7), Inches(0.95), "33", "training rows", BLUE)
    add_metric(slide, Inches(2.75), Inches(1.55), Inches(1.35), Inches(0.95), "28", "retrieval rows", TEAL)
    add_metric(slide, Inches(4.35), Inches(1.55), Inches(1.7), Inches(0.95), "0", "pending evals", AMBER)
    add_metric(slide, Inches(6.3), Inches(1.55), Inches(1.7), Inches(0.95), "11045911", "retry done", GREEN)
    add_text(
        slide,
        Inches(0.9),
        Inches(3.0),
        Inches(5.55),
        Inches(0.6),
        "Final status as of May 7",
        15,
        NAVY,
        True,
    )
    rows = [
        ["11040135", "3-class", "training-strategy retrieval/fusion completed"],
        ["11040134", "8-class", "GigaPath last-2-block warmup/cosine completed"],
        ["11045911", "8-class", "GigaPath ablation retrieval tasks 8-11 completed as chunked retry"],
    ]
    add_table(slide, Inches(0.9), Inches(3.55), Inches(6.4), Inches(1.65), ["Job", "Target", "Purpose"], rows, [Inches(1.0), Inches(1.15), Inches(4.25)], 8.4)
    add_text(
        slide,
        Inches(7.2),
        Inches(3.05),
        Inches(4.9),
        Inches(1.35),
        "Snapshot files:\n/nfs/roberts/project/pi_xy48/hw646/Medical-SAM3/results/visium_hd_exp1/summaries/latest_training_summary.csv\n/nfs/roberts/project/pi_xy48/hw646/Medical-SAM3/results/visium_hd_exp1/summaries/latest_retrieval_summary.csv",
        11,
        NAVY,
    )
    add_text(
        slide,
        Inches(1.0),
        Inches(5.55),
        Inches(11.2),
        Inches(0.55),
        "Current readout: all queued retrieval evals are complete. Last-2 fine-tuning improves classification, fused retrieval, and concat probe; raw R@1 remains strongest for CE-only / last-1.",
        13,
        NAVY,
        False,
        PP_ALIGN.CENTER,
    )


def slide_final_story(prs):
    slide = blank(prs)
    add_title(slide, "Current interpretation and next decision", "The model direction is promising, but the loss design needs ablation before integration into SAM3")
    add_footer(slide)
    bullets = [
        ("What is solid now", "GigaPath crop128 3-class is the most reliable line; CE-only currently has strongest raw R@1, while InfoNCE helps top-k/fusion."),
        ("What improved", "Haiku-inspired last-2-block fine-tuning raised 3-class macro-F1 to 0.694, image accuracy to 0.831, fused R@1 to 0.880, and concat probe to 0.866."),
        ("What 8-class says", "Last-2 warmup/cosine improved 8-class image accuracy to 0.708 and macro-F1 to 0.451, so fine-tuning helps but does not fully solve fine-label noise."),
        ("Next decision", "Use last-2 for fused/probe-style SAM3 region ranking; keep CE-only/last-1 as raw nearest-neighbor retrieval baselines."),
    ]
    for i, (head, body) in enumerate(bullets):
        y = Inches(1.5 + i * 1.25)
        add_label(slide, Inches(0.8), y, Inches(2.4), Inches(0.55), head, [BLUE, AMBER, TEAL, GREEN][i], LIGHT, 11.5)
        add_text(slide, Inches(3.55), y + Inches(0.05), Inches(8.4), Inches(0.5), body, 13, NAVY)


def main() -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    delete_slides_from(prs, 0)
    add_future_directions_overview(prs)
    for fn in [
        slide_ours_what_it_trains_detailed,
        slide_ours_architecture_detailed,
        slide_ours_input_construction_detailed,
        slide_ours_projection_loss_detailed,
        slide_ours_training_strategy_detailed,
        slide_training_variants_map,
        slide_our_experiment_settings,
        slide_completed_runs,
        slide_key_training_result_tables,
        slide_training_curves,
        slide_retrieval_results,
        slide_haiku_what_it_does_detailed,
        slide_haiku_architecture_detailed,
        slide_haiku_input_construction_detailed,
        slide_haiku_projection_loss_detailed,
        slide_haiku_training_fusion_tasks_detailed,
        slide_our_vs_haiku,
        slide_ablation_plan,
        slide_current_hpc_status,
        slide_final_story,
    ]:
        fn(prs)
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    prs.save(FINAL)
    mirror_dir = ROOT / "output" / "group_meeting_ppt"
    mirror_dir.mkdir(parents=True, exist_ok=True)
    prs.save(mirror_dir / FINAL.name)
    print(FINAL)


if __name__ == "__main__":
    main()
